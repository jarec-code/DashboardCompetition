"""Two-step imports with preview (assignment section 6).

1. `preview()` reads the file (CSV, Excel .xlsx, XML feed: Heureka or Google Merchant), normalizes every
   row, validates it and marks it with a status and a reason. Nothing is saved yet. Product rows:
   ok / warning (new), update / unchanged (existing product), review (ambiguous, not saved),
   duplicate (repeated in the file), error.
2. `commit()` checks the rows again against the current catalogue, saves them and reports what happened.

Limits: 5 MB per file, 10 000 rows. XML is parsed with defusedxml (no entities, no external DTDs).
"""
import csv
import html
import io
from datetime import UTC, datetime
from decimal import Decimal
from typing import Any
from urllib.parse import urlsplit

from defusedxml import ElementTree
from pydantic import ValidationError
from sqlalchemy import func, select
from sqlalchemy.orm import Session

from app.models import Competitor, CompetitorProduct, ImportBatch, Product, ProductPrice
from app.models.products import SOURCE_RANK, looks_like_bundle
from app.schemas.products import CURRENCY_MARKET, ProductCreate
from app.scraping.normalize import detect_currency, parse_price
from app.services import matching
from app.services.identifiers import normalize_ean, normalize_isbn, normalize_sku
from app.services.urls import InvalidUrl, normalize_url

MAX_BYTES = 5 * 1024 * 1024
MAX_ROWS = 10_000

# Column names accepted in files (lowercase) -> internal field
ALIASES = {
    "name": "name", "názov": "name", "nazov": "name", "název": "name", "nazev": "name", "productname": "name",
    "product": "name", "title": "name", "price": "price", "cena": "price", "price_vat": "price",
    "market": "market", "trh": "market", "currency": "currency", "mena": "currency", "měna": "currency",
    "isbn": "isbn", "ean": "ean", "gtin": "ean", "sku": "sku", "kód": "sku", "kod": "sku", "item_id": "internal_id",
    "internal_id": "internal_id", "id": "internal_id", "brand": "brand", "manufacturer": "brand",
    "vydavatel": "brand", "vydavateľ": "brand", "category": "category", "categorytext": "category",
    "product_type": "category", "kategória": "category", "kategorie": "category", "priority": "priority",
    "priorita": "priority", "url": "url", "link": "url", "variant": "variant",
    "product_isbn": "product_isbn", "product_ean": "product_ean", "product_sku": "product_sku",
    "product_id": "product_internal_id", "product_name": "product_name", "produkt": "product_name",
    "cena_eshop": "price", "kategorie_qlik": "category", "image_url": "image_url", "image_link": "image_url",
    "imgurl": "image_url",
    "obrázok": "image_url", "obrázek": "image_url", "availability": "availability", "dostupnost": "availability",
    "dostupnosť": "availability", "subject": "attr_subject", "předměty": "attr_subject", "predmety": "attr_subject",
    "grade": "attr_grade", "ročníky": "attr_grade", "ročníky_": "attr_grade", "rocniky": "attr_grade",
    "type": "attr_type", "druh_produktu": "attr_type", "year": "attr_year", "rok_uvedení_na_trh": "attr_year",
}


class ImportFileError(ValueError):
    pass


# ---------- reading files ----------

def _decode(data: bytes) -> str:
    for encoding in ("utf-8-sig", "cp1250"):
        try:
            return data.decode(encoding)
        except UnicodeDecodeError:
            continue
    raise ImportFileError("Súbor nie je v kódovaní UTF-8 ani Windows-1250.")


def _map_header(names: list[str]) -> list[str | None]:
    mapped = [ALIASES.get((n or "").strip().lower().replace(" ", "_")) for n in names]
    if not any(mapped):
        raise ImportFileError("Hlavička neobsahuje žiadny známy stĺpec (napr. name, price, url).")
    return mapped


def _rows_from_table(header: list[str], records) -> list[dict[str, Any]]:
    mapped = _map_header(header)
    rows = []
    for number, record in records:
        values = {field: str(value).strip() for field, value in zip(mapped, record, strict=False)  # short rows are fine
                  if field and value is not None and str(value).strip()}
        if values:
            rows.append({"row": number, **values})
        if len(rows) > MAX_ROWS:
            raise ImportFileError(f"Súbor má viac ako {MAX_ROWS} riadkov.")
    return rows


def _read_csv(data: bytes) -> list[dict[str, Any]]:
    text = _decode(data)
    if not text.strip():
        raise ImportFileError("Súbor je prázdny.")
    try:
        dialect = csv.Sniffer().sniff(text.splitlines()[0], delimiters=",;\t")
    except csv.Error:
        dialect = csv.excel
    reader = csv.reader(io.StringIO(text), dialect=dialect)
    header = next(reader)
    return _rows_from_table(header, ((i, rec) for i, rec in enumerate(reader, start=2)))


def _read_xlsx(data: bytes) -> list[dict[str, Any]]:
    from openpyxl import load_workbook

    try:
        workbook = load_workbook(io.BytesIO(data), read_only=True, data_only=True)
    except Exception:
        raise ImportFileError("Súbor sa nedá otvoriť ako Excel (.xlsx).") from None
    sheet = workbook.worksheets[0]
    rows = sheet.iter_rows(values_only=True)
    header = [str(c) if c is not None else "" for c in next(rows, [])]
    result = _rows_from_table(header, ((i, rec) for i, rec in enumerate(rows, start=2)))
    workbook.close()
    return result


def _text(element, path: str, namespaces: dict | None = None) -> str | None:
    found = element.find(path, namespaces or {})
    return found.text.strip() if found is not None and found.text and found.text.strip() else None


def _availability_from_delivery(item) -> str | None:
    """Heureka DELIVERY_DATE: 0 = in stock, a number of days or a date = not in stock now."""
    value = _text(item, "DELIVERY_DATE")
    if value is None:
        return None
    return "in_stock" if value == "0" else "out_of_stock"


def _read_xml(data: bytes) -> tuple[str, list[dict[str, Any]]]:
    try:
        root = ElementTree.fromstring(data, forbid_dtd=True)
    except Exception as exc:
        raise ImportFileError(f"XML sa nedá spracovať: {exc.__class__.__name__}.") from None
    rows: list[dict[str, Any]] = []
    if root.tag.upper() == "SHOP":  # Heureka feed
        for i, item in enumerate(root.iter("SHOPITEM"), start=1):
            # EAN may repeat (e.g. an invalid code first, the valid one second): prefer a valid code.
            eans = [e.text.strip() for e in item.findall("EAN") if e.text and e.text.strip()]
            eans.sort(key=lambda code: normalize_ean(code) is None)
            row = {"row": i, "name": _text(item, "PRODUCTNAME") or _text(item, "PRODUCT"),
                   "price": _text(item, "PRICE_VAT"), "ean": eans[0] if eans else None, "isbn": _text(item, "ISBN"),
                   "internal_id": _text(item, "ITEM_ID"), "url": _text(item, "URL"),
                   "image_url": _text(item, "IMGURL"), "category": _text(item, "CATEGORYTEXT"),
                   "brand": _text(item, "MANUFACTURER"), "availability": _availability_from_delivery(item)}
            rows.append({k: v for k, v in row.items() if v is not None})
        fmt = "heureka"
    elif root.tag in ("rss", "feed") or root.tag.endswith("}feed"):  # Google Merchant (RSS or Atom)
        ns = {"g": "http://base.google.com/ns/1.0", "a": "http://www.w3.org/2005/Atom"}
        items = list(root.iter("item")) or list(root.iter("{http://www.w3.org/2005/Atom}entry"))
        for i, item in enumerate(items, start=1):
            price = _text(item, "g:price", ns) or ""
            row = {"row": i, "name": _text(item, "title") or _text(item, "g:title", ns) or _text(item, "a:title", ns),
                   "price": price, "currency": detect_currency(price), "ean": _text(item, "g:gtin", ns),
                   "internal_id": _text(item, "g:id", ns), "url": _text(item, "link") or _text(item, "g:link", ns),
                   "category": _text(item, "g:product_type", ns), "brand": _text(item, "g:brand", ns),
                   "image_url": _text(item, "g:image_link", ns), "availability": _text(item, "g:availability", ns)}
            rows.append({k: v for k, v in row.items() if v is not None})
        fmt = "google"
    else:
        raise ImportFileError("Neznámy XML formát. Podporované sú feedy Heureka (SHOP/SHOPITEM) a Google Merchant.")
    if len(rows) > MAX_ROWS:
        raise ImportFileError(f"Feed má viac ako {MAX_ROWS} položiek.")
    return fmt, rows


def read_file(data: bytes, file_name: str) -> tuple[str, list[dict[str, Any]]]:
    if len(data) > MAX_BYTES:
        raise ImportFileError("Súbor je väčší ako 5 MB.")
    if not data.strip():
        raise ImportFileError("Súbor je prázdny.")
    name = file_name.lower()
    if name.endswith(".xlsx") or data[:2] == b"PK":
        return "xlsx", _read_xlsx(data)
    if name.endswith(".xml") or data.lstrip()[:1] == b"<":
        return _read_xml(data)
    return "csv", _read_csv(data)


# ---------- products ----------
#
# Each row is matched to an existing product, in this order: internal ID, SKU, ISBN/EAN, product URL, and
# finally the name together with the same price (a name alone is never enough: editions share names).
# A match updates the product (identifiers are only filled in, never replaced; other fields are replaced
# only by a source at least as trusted as the one that set them, see SOURCE_RANK). Ambiguous rows go to
# manual review and are not saved. Re-running the same import changes nothing (idempotent).

def _first_error(exc: ValidationError) -> str:
    err = exc.errors()[0]
    field = ".".join(str(p) for p in err["loc"] if p != "prices" and not isinstance(p, int))
    return f"{field}: {err['msg']}" if field else err["msg"]


STRONG_KEYS = ("internal_id", "sku", "gtin", "url")
KEY_LABELS = {"internal_id": "interné ID", "sku": "SKU", "gtin": "ISBN/EAN", "url": "URL", "isbn": "ISBN",
              "ean": "EAN"}
DESCRIPTIVE = ("name", "brand", "category", "url", "image_url")
FIELD_LABELS = {"name": "názov", "brand": "značka", "category": "kategória", "url": "URL", "image_url": "obrázok",
                "availability": "dostupnosť", "attributes": "parametre", "is_bundle": "balíček",
                "internal_id": "interné ID", "sku": "SKU", "isbn": "ISBN", "ean": "EAN", "price": "cena"}
FEED_FORMATS = ("heureka", "google")
_AVAILABILITY = (
    ("in_stock", ("in stock", "instock", "in_stock", "skladem", "na sklade", "available", "dostupn")),
    ("out_of_stock", ("out of stock", "out_of_stock", "vyprodáno", "vypredané", "nedostupn", "sold out")),
    ("preorder", ("preorder", "pre-order", "předprodej", "predpredaj", "připravujeme")),
)


def _availability(value: str | None) -> str | None:
    text = (value or "").strip().lower()
    if not text:
        return None
    for code, words in _AVAILABILITY:
        if any(text.startswith(w) for w in words):
            return code
    return "unknown"


def _name_key(name: str) -> str:
    return matching.normalize_name(html.unescape(name))


def _gtins(data: dict[str, Any]) -> list[str]:
    return [code for code in (data.get("isbn"), data.get("ean")) if code]


def _identifiers(row: dict[str, Any], notes: list[str]) -> dict[str, str | None]:
    """Valid identifiers of the row; invalid ones are dropped with a note instead of rejecting the row."""
    result: dict[str, str | None] = {}
    for field, normalize in (("isbn", normalize_isbn), ("ean", normalize_ean)):
        raw = (row.get(field) or "").strip()
        value = normalize(raw) if raw else None
        if raw and value is None:
            notes.append(f"Neplatné {KEY_LABELS[field]} „{raw}“ sa ignoruje.")
        result[field] = value
    if result["ean"] and result["ean"][:3] in ("978", "979") and normalize_isbn(result["ean"]):
        # An EAN starting 978/979 is the ISBN-13 of a book: store it once, as ISBN.
        if not result["isbn"]:
            result["isbn"] = result["ean"]
        if result["isbn"] == result["ean"]:
            result["ean"] = None
    result["sku"] = row.get("sku")
    result["internal_id"] = row.get("internal_id")
    return result


def _validate_product(row: dict[str, Any], default_market: str | None) -> tuple[dict[str, Any], list[str]]:
    notes: list[str] = []
    raw_price = row.get("price")
    price = parse_price(raw_price)
    if raw_price and price is None:
        raise ValueError(f"Neplatná cena: „{raw_price}“.")
    currency = (row.get("currency") or "").upper() or detect_currency(raw_price)
    market = (row.get("market") or "").upper() or CURRENCY_MARKET.get(currency or "") or default_market
    prices = []
    if price is not None and price > 0:
        if not market:
            raise ValueError("Chýba trh (stĺpec market), známa mena alebo predvolený trh importu.")
        prices = [{"market": market, "price": price, "currency": currency or None}]
    else:
        notes.append("Cena v zdroji je 0, produkt bude bez ceny." if price is not None else "Chýba cena.")
    attributes = {key: row.get(f"attr_{key}") for key in ("subject", "grade", "type", "year")}
    try:
        product = ProductCreate(
            name=html.unescape(row.get("name", "")), **_identifiers(row, notes), brand=row.get("brand"),
            category=(row.get("category") or "")[:200] or None, priority=row.get("priority") or "regular",
            url=row.get("url"), image_url=row.get("image_url"), availability=_availability(row.get("availability")),
            attributes={k: v[:200] for k, v in attributes.items() if v}, prices=prices)
    except ValidationError as exc:
        raise ValueError(_first_error(exc)) from None
    data = product.model_dump(mode="json")
    data["attributes"] = {k: v for k, v in (data["attributes"] or {}).items() if v}
    data["is_bundle"] = True if looks_like_bundle(data["name"]) else None
    return data, notes


class _Catalogue:
    """Our products indexed by every matching key, loaded once per preview or commit."""

    def __init__(self, db: Session):
        self.by_key: dict[str, dict[str, Product]] = {key: {} for key in STRONG_KEYS}
        self.by_name: dict[str, list[Product]] = {}
        for product in db.scalars(select(Product)):
            self.add(product)

    def add(self, product: Product) -> None:
        values = {"internal_id": [product.internal_id], "sku": [product.sku], "url": [product.url],
                  "gtin": [product.isbn, product.ean]}
        for key, codes in values.items():
            for code in codes:
                if code:
                    self.by_key[key][code] = product
        bucket = self.by_name.setdefault(_name_key(product.name), [])
        if product not in bucket:
            bucket.append(product)

    def owner(self, key: str, value: str | None) -> Product | None:
        return self.by_key[key].get(value) if value else None

    def resolve(self, data: dict[str, Any]) -> tuple[str, Product | None, str | None]:
        """("new" | "match" | "review", product, note)."""
        found: dict[int, tuple[Product, list[str]]] = {}
        for key in STRONG_KEYS:
            for value in (_gtins(data) if key == "gtin" else [data.get(key)]):
                product = self.owner(key, value)
                if product and key not in found.setdefault(product.id, (product, []))[1]:
                    found[product.id][1].append(key)
        if len(found) > 1:
            listed = "; ".join(f"„{p.name}“ (#{p.id}, {', '.join(KEY_LABELS[k] for k in keys)})"
                               for p, keys in found.values())
            return "review", None, f"Identifikátory riadku patria rôznym produktom: {listed}. Skontrolujte ich ručne."
        if found:
            return "match", next(iter(found.values()))[0], None
        same_name = [p for p in self.by_name.get(_name_key(data["name"]), []) if not _conflicts(p, data)]
        if not same_name:
            return "new", None, None
        if len(same_name) > 1:
            listed = ", ".join(f"#{p.id}" for p in same_name)
            return "review", None, (f"Rovnaký názov má viac produktov ({listed}); bez identifikátora ich nemožno "
                                    "rozlíšiť.")
        product = same_name[0]
        if data["prices"] and _same_price(product, data["prices"][0]):
            return "match", product, "Spárované podľa názvu a rovnakej ceny."
        return "review", None, (f"Možná zhoda s produktom „{product.name}“ (#{product.id}) iba podľa názvu. Ak ide "
                                "o ten istý produkt, doplňte mu interné ID alebo ISBN a import zopakujte.")


def _conflicts(product: Product, data: dict[str, Any]) -> bool:
    """True when the row and the product carry different values of the same identifier (another edition)."""
    if data.get("internal_id") and product.internal_id and data["internal_id"] != product.internal_id:
        return True
    if data.get("sku") and product.sku and data["sku"] != product.sku:
        return True
    row_codes, own_codes = set(_gtins(data)), {c for c in (product.isbn, product.ean) if c}
    if row_codes and own_codes and not row_codes & own_codes:
        return True
    return bool(data.get("url") and product.url and data["url"] != product.url)


def _same_price(product: Product, price: dict[str, Any]) -> bool:
    own = next((p for p in product.prices if p.market == price["market"]), None)
    return own is not None and own.price == Decimal(price["price"]) and own.currency == price["currency"]


def _changes(product: Product, data: dict[str, Any], rank: int, catalogue: _Catalogue
             ) -> tuple[dict[str, Any], list[str]]:
    """Field -> new value for an existing product, and notes about values that were kept."""
    may_replace = rank >= SOURCE_RANK.get(product.source_kind, 0)
    changes: dict[str, Any] = {}
    notes: list[str] = []
    for field in ("internal_id", "sku", "isbn", "ean"):
        new, old = data.get(field), getattr(product, field)
        if not new or new == old:
            continue
        if field in ("isbn", "ean") and new in (product.isbn, product.ean):
            continue  # the same code stored in the other column
        key = "gtin" if field in ("isbn", "ean") else field
        if old:
            notes.append(f"{KEY_LABELS[field]} v zdroji ({new}) sa líši od uloženého ({old}), ponechané uložené.")
        elif catalogue.owner(key, new) not in (None, product):
            notes.append(f"{KEY_LABELS[field]} {new} má iný produkt, nedoplnené.")
        else:
            changes[field] = new
    for field in DESCRIPTIVE:
        new, old = data.get(field), getattr(product, field)
        if new and new != old and (not old or may_replace):
            if field == "url" and catalogue.owner("url", new) not in (None, product):
                notes.append(f"URL {new} má iný produkt, nezmenené.")
                continue
            changes[field] = new
    availability = data.get("availability")
    if availability and availability != "unknown" and availability != product.availability \
            and (product.availability == "unknown" or may_replace):
        changes["availability"] = availability
    if data.get("is_bundle") and product.is_bundle is None:  # a decision made by hand is kept
        changes["is_bundle"] = True
    attributes = dict(product.attributes or {})
    for key, value in (data.get("attributes") or {}).items():
        if value and attributes.get(key) != value and (not attributes.get(key) or may_replace):
            attributes[key] = value
    if attributes != (product.attributes or {}):
        changes["attributes"] = attributes
    for price in data["prices"]:
        own = next((p for p in product.prices if p.market == price["market"]), None)
        if own is None or (may_replace and not _same_price(product, price)):
            changes["price"] = price
    return changes, notes


def _describe_changes(product: Product, changes: dict[str, Any]) -> str:
    parts = []
    for field, value in changes.items():
        if field == "price":
            own = next((p for p in product.prices if p.market == value["market"]), None)
            before = f"{own.price} {own.currency} → " if own else ""
            parts.append(f"cena {value['market']} {before}{value['price']} {value['currency']}")
        else:
            parts.append(FIELD_LABELS[field])
        if field == "name":
            parts[-1] = f"názov „{product.name}“ → „{value}“"
    return "Zmení sa: " + ", ".join(parts) + "."


def _row_keys(data: dict[str, Any]) -> list[str]:
    keys = [f"{k}:{data[k]}" for k in ("internal_id", "sku", "url") if data.get(k)]
    keys += [f"gtin:{code}" for code in _gtins(data)]
    return keys or [f"name:{_name_key(data['name'])}"]


def _check_product_row(catalogue: _Catalogue, data: dict[str, Any], notes: list[str], rank: int
                       ) -> dict[str, Any]:
    """Status of one valid row against the catalogue (shared by preview and commit)."""
    outcome, product, note = catalogue.resolve(data)
    notes = notes + ([note] if note else [])
    if outcome == "review":
        return {"status": "review", "reason": " ".join(notes)}
    if outcome == "new":
        if not data["prices"] and not data.get("url"):
            return {"status": "error", "reason": "Bez ceny a bez URL: nejde o overiteľný predajný produkt."}
        return {"status": "warning" if notes else "ok", "reason": " ".join(notes) or None}
    changes, change_notes = _changes(product, data, rank, catalogue)
    notes += change_notes
    if not changes:
        return {"status": "unchanged", "product_id": product.id,
                "reason": " ".join([f"Bez zmeny (produkt #{product.id})."] + notes)}
    return {"status": "update", "product_id": product.id,
            "reason": " ".join([f"Produkt #{product.id}. {_describe_changes(product, changes)}"] + notes)}


def _source(batch_format: str, file_name: str) -> tuple[str, str]:
    """(source_kind, label) of a batch: XML feeds of an e-shop are trusted more than uploaded tables."""
    kind = "feed" if batch_format in FEED_FORMATS else "file"
    host = urlsplit(file_name).hostname if file_name.startswith(("http://", "https://")) else None
    return kind, (f"{host.removeprefix('www.')} (feed)" if host else file_name)[:255]


def _preview_products(db: Session, rows: list[dict[str, Any]], file_format: str,
                      default_market: str | None) -> list[dict[str, Any]]:
    catalogue, rank = _Catalogue(db), SOURCE_RANK[_source(file_format, "")[0]]
    result, seen = [], {}
    for row in rows:
        try:
            data, notes = _validate_product(row, default_market)
        except ValueError as exc:
            result.append({"row": row["row"], "status": "error", "reason": str(exc), "data": row})
            continue
        clash = next((seen[k] for k in _row_keys(data) if k in seen), None)
        if clash:
            result.append({"row": row["row"], "status": "duplicate",
                           "reason": f"Rovnaký produkt je už v riadku {clash}.", "data": data})
            continue
        checked = _check_product_row(catalogue, data, notes, rank)
        result.append({"row": row["row"], **checked, "data": data, "notes": notes})
        if checked["status"] != "error":
            for k in _row_keys(data):
                seen[k] = row["row"]
    return result


def _apply(product: Product, changes: dict[str, Any]) -> None:
    for field, value in changes.items():
        if field == "price":
            own = next((p for p in product.prices if p.market == value["market"]), None)
            if own is None:
                product.prices.append(ProductPrice(market=value["market"], price=Decimal(value["price"]),
                                                   currency=value["currency"], vat_mode=value["vat_mode"]))
            else:
                own.price, own.currency, own.vat_mode = Decimal(value["price"]), value["currency"], value["vat_mode"]
        else:
            setattr(product, field, value)


def _commit_products(db: Session, rows: list[dict[str, Any]], batch: ImportBatch) -> dict[str, Any]:
    kind, label = _source(batch.file_format, batch.file_name)
    rank, now = SOURCE_RANK[kind], datetime.now(UTC)
    catalogue = _Catalogue(db)
    counts = {"imported": 0, "updated": 0, "unchanged": 0}
    skipped = []
    for row in rows:
        data = row["data"]
        # Checked again: the catalogue may have changed since the preview.
        checked = _check_product_row(catalogue, data, row.get("notes") or [], rank)
        if checked["status"] in ("review", "error"):
            skipped.append({"row": row["row"], "reason": checked["reason"]})
            continue
        if checked["status"] in ("ok", "warning"):
            product = Product(**{k: v for k, v in data.items() if k != "prices"})
            product.prices = [ProductPrice(market=p["market"], price=Decimal(p["price"]), currency=p["currency"],
                                           vat_mode=p["vat_mode"]) for p in data["prices"]]
            product.source_kind, product.source, product.source_synced_at = kind, label, now
            db.add(product)
            db.flush()
            counts["imported"] += 1
        else:
            product = db.get(Product, checked["product_id"])
            changes, _ = _changes(product, data, rank, catalogue)
            _apply(product, changes)
            if rank >= SOURCE_RANK.get(product.source_kind, 0):
                product.source_kind, product.source, product.source_synced_at = kind, label, now
            db.flush()
            counts["updated" if changes else "unchanged"] += 1
        catalogue.add(product)
    if counts["imported"] or counts["updated"]:
        matching.rematch_open_offers(db)
    return {**counts, "skipped_at_commit": skipped}


# ---------- offers ----------

def _competitor_for(competitors: list[Competitor], url: str) -> Competitor | None:
    host = urlsplit(url).hostname or ""
    for c in competitors:
        if host == c.domain or host.endswith("." + c.domain):
            return c
    return None


def _find_product(db: Session, row: dict[str, Any]) -> tuple[Product | None, str | None]:
    """Product referenced by the row (identifiers first, then the exact name). Returns (product, problem)."""
    refs = {"isbn": normalize_isbn(row.get("product_isbn")), "ean": normalize_ean(row.get("product_ean")),
            "sku": normalize_sku(row.get("product_sku")), "internal_id": row.get("product_internal_id")}
    for field, value in refs.items():
        if value:
            product = db.scalar(select(Product).where(getattr(Product, field) == value))
            if product:
                return product, None
            return None, f"Produkt s {field.upper()} {value} neexistuje, ponuka sa pridá nepriradená."
    if row.get("product_name"):
        product = db.scalars(select(Product).where(func.lower(Product.name) == row["product_name"].lower())).first()
        return (product, None) if product else (None, "Produkt s týmto názvom neexistuje, ponuka sa pridá "
                                                      "nepriradená.")
    return None, None


def _preview_offers(db: Session, rows: list[dict[str, Any]], competitor_id: int | None) -> list[dict[str, Any]]:
    competitors = list(db.scalars(select(Competitor).where(Competitor.status != "archived")))
    if competitor_id is not None:
        competitors = [c for c in competitors if c.id == competitor_id]
        if not competitors:
            raise ImportFileError("Vybraný konkurent neexistuje alebo je archivovaný.")
    result, seen = [], {}
    for row in rows:
        raw = row.get("url")
        if not raw:
            result.append({"row": row["row"], "status": "error", "reason": "Chýba URL.", "data": row})
            continue
        try:
            url = normalize_url(raw)
        except InvalidUrl as exc:
            result.append({"row": row["row"], "status": "error", "reason": str(exc), "data": row})
            continue
        competitor = _competitor_for(competitors, url)
        if competitor is None:
            result.append({"row": row["row"], "status": "error", "data": row,
                           "reason": "URL nepatrí žiadnemu evidovanému konkurentovi (ani vybranému)."})
            continue
        variant = (row.get("variant") or "").strip()[:100]
        product, problem = _find_product(db, row)
        data = {"url": url, "variant": variant, "competitor_id": competitor.id, "competitor_name": competitor.name,
                "product_id": product.id if product else None, "product_name": product.name if product else None}
        key = (competitor.id, url, variant)
        exists = db.scalar(select(CompetitorProduct.id).where(
            CompetitorProduct.competitor_id == competitor.id, CompetitorProduct.product_url == url,
            CompetitorProduct.variant == variant))
        if exists:
            result.append({"row": row["row"], "status": "duplicate", "reason": "Táto URL je už sledovaná.",
                           "data": data})
        elif key in seen:
            result.append({"row": row["row"], "status": "duplicate",
                           "reason": f"Rovnaká URL je už v riadku {seen[key]}", "data": data})
        else:
            seen[key] = row["row"]
            result.append({"row": row["row"], "status": "warning" if problem else "ok", "reason": problem,
                           "data": data})
    return result


def _commit_offers(db: Session, rows: list[dict[str, Any]]) -> tuple[int, list[dict]]:
    imported, skipped = 0, []
    for row in rows:
        data = row["data"]
        exists = db.scalar(select(CompetitorProduct.id).where(
            CompetitorProduct.competitor_id == data["competitor_id"], CompetitorProduct.product_url == data["url"],
            CompetitorProduct.variant == data["variant"]))
        if exists:
            skipped.append({"row": row["row"], "reason": "URL medzitým pribudla."})
            continue
        offer = CompetitorProduct(competitor_id=data["competitor_id"], product_url=data["url"], variant=data["variant"])
        db.add(offer)
        db.flush()
        if data.get("product_id") and db.get(Product, data["product_id"]):
            matching.assign_manually(offer, data["product_id"])
        imported += 1
    return imported, skipped


# ---------- public API ----------

STATUSES = ("ok", "warning", "update", "unchanged", "review", "duplicate", "error")


def _stats(rows: list[dict[str, Any]]) -> dict[str, int]:
    return {s: sum(1 for r in rows if r["status"] == s) for s in STATUSES}


def preview(db: Session, kind: str, data: bytes, file_name: str, user_id: int | None,
            competitor_id: int | None = None, default_market: str | None = None) -> ImportBatch:
    file_format, rows = read_file(data, file_name)
    if not rows:
        raise ImportFileError("Súbor neobsahuje žiadne riadky s údajmi.")
    if kind == "products":
        checked = _preview_products(db, rows, file_format, (default_market or "").strip().upper() or None)
    else:
        checked = _preview_offers(db, rows, competitor_id)
    batch = ImportBatch(kind=kind, status="preview", file_name=file_name[:255], file_format=file_format,
                        competitor_id=competitor_id, created_by=user_id, rows=checked, stats=_stats(checked))
    db.add(batch)
    db.flush()
    return batch


SAVED_STATUSES = ("ok", "warning", "update", "unchanged")


def commit(db: Session, batch: ImportBatch) -> dict[str, Any]:
    if batch.status != "preview":
        raise ImportFileError("Tento import už bol potvrdený alebo zrušený.")
    to_save = [r for r in batch.rows if r["status"] in SAVED_STATUSES]
    if batch.kind == "products":
        outcome = _commit_products(db, to_save, batch)
    else:
        imported, skipped = _commit_offers(db, to_save)
        outcome = {"imported": imported, "updated": 0, "unchanged": 0, "skipped_at_commit": skipped}
    skipped = outcome["skipped_at_commit"]
    batch.status, batch.committed_at = "committed", datetime.now(UTC)
    batch.stats = {**batch.stats, "imported": outcome["imported"], "updated": outcome["updated"],
                   "unchanged_saved": outcome["unchanged"], "skipped_at_commit": len(skipped)}
    return {"processed": len(batch.rows), "imported": outcome["imported"], "updated": outcome["updated"],
            "unchanged": outcome["unchanged"], "review": batch.stats.get("review", 0),
            "duplicates": batch.stats["duplicate"], "errors": batch.stats["error"], "skipped_at_commit": skipped}
